import csv
import logging
from datetime import datetime

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

import config
from Utilization import calculate_utilization, read_samples

path = config.EXCEL_PATH  # Use the EXCEL_PATH from config.py

logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)

def write_to_excel(utilization_data, path, date):

    if path.exists():         #Check if the Excel file already exists
        wb = load_workbook(path)    #Load the existing Excel file
        ws = wb.active

    else:
        wb = Workbook()        #Create a new Excel workbook if a current file does not already exist
        ws = wb.active
        ws.title = "Utilization Data"

        # Write headers to the first row of the Excel sheet
        headers = ["Date","Powered On Time (hrs)", "Active Time (hrs)", "Idle Time (hrs)", "Utilization (%)"]      
        ws.append(headers)
        header_fill = PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid")
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")            


    # Write data to the next available row in the Excel sheet
    data_row = [
        date,
        round(utilization_data.powered_on_time/3600,1),  # Convert seconds to hours and round to 1 decimal place
        round(utilization_data.active_time/3600,1),  # Convert seconds to hours and round to 1 decimal place
        round(utilization_data.idle_time/3600,1),  # Convert seconds to hours and round to 1 decimal place
        round(utilization_data.utilization,1) if utilization_data.utilization is not None else "N/A"
    ]
    ws.append(data_row)  #Append data row to the Excel sheet

    thin = Side(style="thin", color="000000")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for row in ws.iter_rows():
        for cell in row:
            cell.border = border
            if cell.row != 1:
                cell.alignment = Alignment(horizontal="center", vertical="center")

    for col_cells in ws.columns:
        max_length = max((len(str(cell.value)) for cell in col_cells if cell.value is not None), default=0)
        col_letter = get_column_letter(col_cells[0].column)
        ws.column_dimensions[col_letter].width = max_length + 4

    # Save the workbook
    wb.save(path)   

def verify_write(path, date):
    try:
        wb = load_workbook(path)                         #Open the Excel file
        ws = wb.active
        last_row = ws.max_row
        last_date = ws.cell(row=last_row, column=1).value          #Get the last date from the first column of the last row
        return last_date == date          #Verify if the last date in the Excel file matches the current date
    except Exception:  # noqa: BLE001
        return False                        #Return False if there is an error while loading the Excel file


def wipe_CSV(CSV_path, Excel_path, date):
    if verify_write(Excel_path, date):
        with open(CSV_path, 'w', newline='') as file:                  # Wipe the CSV file by opening it in write mode
            writer = csv.writer(file)
            writer.writerow(['Timestamp', 'Robot Mode', 'Active Time'])  # Write only the header
        logger.info("CSV file has been wiped.")                  #Print confirmation of CSV wipe

    else:
        logger.warning("CSV file has NOT been wiped. The last entry in the Excel file does not match the current date.")       #Print warning if CSV file has not been wiped



if __name__ == "__main__":
    
    sample = read_samples(config.CSV_PATH)  # Read samples from the CSV file

    utilization_data = calculate_utilization(sample, config.ROBOTMODE_THRESHOLD, config.WRITING_FREQUENCY)  # Calculate utilization data

    date = datetime.now().strftime("%B %Y")   # "May 2026" in local time  # noqa: DTZ005

    try:

        write_to_excel(utilization_data, path, date)  # Write the data to the Excel file

    except Exception as e:  # noqa: BLE001

        logger.error(f"Failed to write to Excel file: {e}. CSV will not be wiped.")  # Log the error if writing to Excel fails

    else:
        wipe_CSV(config.CSV_PATH, config.EXCEL_PATH, date)  # Wipe the CSV file if the write operation was successful




